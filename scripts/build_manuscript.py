#!/usr/bin/env python3
"""Build the final paper-style Word manuscript from validated results."""

import csv
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(os.environ.get("DARK_PROJECT_ROOT", Path(__file__).resolve().parents[1])).resolve()
REVISION_PLOTS = Path(os.environ.get("DARK_REVISION_PLOTS", ROOT / "plots" / "revision")).resolve()
OUT = Path(os.environ.get("DARK_MANUSCRIPT_OUT", ROOT / "paper" / "ML Paper.docx")).resolve()
BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 98, 108)
LIGHT = "E8EEF5"
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)


def set_font(run, name="Cambria", size=10, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_table(doc, headers, rows, widths, align=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), name="Arial", size=8.5, bold=True, color=BLUE)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row_data in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if align and align[i] == "left" else WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(str(text)), name="Arial", size=8.5)
    set_table_geometry(table, widths)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    return table


def add_caption(doc, label, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run(f"{label}. "), name="Arial", size=9, bold=True)
    set_font(p.add_run(text), name="Arial", size=9)


def add_figure(doc, path, number, caption, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("title", f"Figure {number}")
    picture._inline.docPr.set("descr", caption)
    add_caption(doc, f"Figure {number}", caption)


def paragraph(doc, text="", bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), bold=True)
        set_font(p.add_run(text[len(bold_lead):]))
    else:
        set_font(p.add_run(text))
    return p


def equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run(text), name="Cambria Math", size=10.5, italic=True)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_font(p.add_run(text), name="Arial", size={1: 14, 2: 11.5, 3: 10.5}[level], bold=True, color=BLUE)
    return p


def metrics():
    with (ROOT / "results" / "ml" / "ml_metrics.csv").open() as handle:
        return list(csv.DictReader(handle))


def selected(rows, kind, masses):
    lookup = {(r["signal_type"], float(r["mass_GeV"])): r for r in rows}
    output = []
    for mass in masses:
        r = lookup[(kind, mass)]
        output.append([
            f"{mass:g}", f"{float(r['auc']):.3f}",
            f"{float(r['signal_efficiency']):.3f}",
            f"{float(r['photon_background_efficiency']):.3f}",
            f"{float(r['significance_improvement']):.3f} +/- {float(r['significance_improvement_mc_error']):.3f}",
            f"{100*float(r['coupling_reach_improvement']):.1f}%",
        ])
    return output


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Cambria"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after in (("Heading 1", 14, 12, 5), ("Heading 2", 11.5, 9, 4), ("Heading 3", 10.5, 7, 3)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    caption_style = doc.styles["Caption"]
    caption_style.font.name = "Arial"
    caption_style.font.size = Pt(9)
    caption_style.font.italic = False
    caption_style.font.color.rgb = BLACK

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("Nuclear-recoil information for invisible dark-boson selection at the EIC"), name="Arial", size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run("Article"), name="Arial", size=9, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run("When Does Nuclear-Recoil Information Make Machine Learning Useful for Invisible Dark-Boson Selection at the Electron-Ion Collider?"), name="Arial", size=20, bold=True, color=RGBColor(11, 37, 69))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run("Rojae Mighty"), name="Arial", size=12, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run("A controlled parton-level value-of-information study | 27 August 2026"), name="Arial", size=9.5, italic=True, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run("Abstract. "), name="Arial", size=9.5, bold=True, color=BLUE)
    abstract = (
        "We test a specific question: does a more flexible machine-learning selector improve an invisible-dark-boson search when it receives the same information as optimized rectangular cuts, and does the answer change when nuclear-recoil information is added? We implement the coherent exclusive process eAu -> eAu phi at 18 GeV x 100 GeV per nucleon for invisibly decaying scalar and vector bosons. This is electron-line dark-boson bremsstrahlung from an intact gold nucleus, not exclusive vector-meson production. We present the inclusive and selected scalar and vector signal cross sections across the boson-mass range. We then compare boosted decision trees with a strong multistart rectangular selector. Both methods use identical weighted samples, input variables, training/validation/test roles, and optimization objectives. With the recoil-electron variables Q², pT,e, eta_e, and E_e alone, no point in the complete 22-point scalar/vector mass scan shows familywise-controlled evidence of machine-learning superiority, and ten independent generation-and-training replicas confirm practical equivalence at four benchmarks. Adding the exact generator-level nuclear momentum transfer t to both methods changes the high-mass result. On the sealed 10 GeV samples, the weighted signal-versus-photon-proxy AUC rises from 0.674 to 0.822 for the vector and from 0.721 to 0.839 for the scalar. In ten fresh replicas at 10 GeV, the median BDT coupling-threshold advantages over same-input optimized cuts are 1.208% for the vector and 1.361% for the scalar, with hierarchical 95% intervals of [1.081%, 1.326%] and [1.206%, 1.529%]. The BDT wins all ten replicas for both signals, and both familywise lower bounds exceed the predefined 1% practical threshold. The supported conclusion is therefore conditional but specific: classifier complexity alone provides no useful gain from the available electron kinematics, whereas additional independent recoil information creates a small, statistically robust high-mass BDT advantage within our parton-level signal and photon-proxy model. Exact t is an oracle input, not a reconstructed detector observable. The study does not yet establish detector-level EIC sensitivity because intact-ion acceptance and resolution, validated coherent-photon and event-level DIS samples, detector response, systematic uncertainties, and likelihood-based limits are absent."
    )
    set_font(p.add_run(abstract), name="Arial", size=9.5)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run("Keywords: "), name="Arial", size=9, bold=True)
    set_font(p.add_run("Electron-Ion Collider; invisible dark bosons; nuclear recoil; momentum transfer; boosted decision trees; event selection"), name="Arial", size=9)

    heading(doc, "1. Introduction", 1)
    paragraph(doc, "Invisible particles do not leave ordinary tracks or calorimeter deposits, so they must be inferred from the visible particles that recoil against them and from an imbalance in energy or momentum. Davoudiasl and Liu proposed that the Electron-Ion Collider (EIC) could search for an invisibly decaying scalar or vector boson phi through the coherent exclusive process e Au -> e Au phi [1,2]. The word exclusive means that the complete final state is specified and the gold nucleus remains intact. Coherence over the nuclear charge supplies an approximate Z-squared enhancement, while the recoil electron and the outgoing intact ion carry complementary information about the event.")
    paragraph(doc, "The reference search uses mass-dependent rectangular requirements on four recoil-electron quantities: the electron momentum transfer Q², the electron transverse momentum pT,e, the electron pseudorapidity eta_e, and the electron energy E_e. The nuclear momentum transfer enters the underlying cross-section calculation and nuclear form factor, but it is not supplied to the published event selection as an additional discriminating variable. We denote the positive squared momentum transferred to the nucleus by t. If the outgoing ion can be reconstructed, t provides information from the ion side of the event that is not completely determined by the measured recoil electron when phi is invisible.")
    paragraph(doc, "A machine-learning model can draw a more complicated decision boundary than independent upper and lower cuts. That additional flexibility is useful only when the input variables contain correlations that a rectangular selector cannot already capture. A comparison with the published Table I requirements alone would not isolate this question because those cuts were optimized in a different simulation and do not receive every variable considered by the BDT. We therefore compare the BDT with a strong multistart rectangular selector optimized on the same weighted samples and the same objective. Within each comparison, the two methods always receive exactly the same information.")

    heading(doc, "1.1. Questions addressed and claims tested", 2)
    paragraph(doc, "The analysis is organized around four increasingly demanding questions. First, how do our inclusive and selected scalar and vector signal cross sections depend on the boson mass? Second, when both algorithms receive only the recoil-electron quantities used by the reference search, does the BDT outperform a reoptimized rectangular selector? Third, when exact generator-level t is added to both algorithms, does a repeatable high-mass BDT advantage appear? Fourth, what additional detector and background work is required before that generator-level result can be interpreted as an EIC sensitivity projection?")
    paragraph(doc, "The principal methodological contribution is a controlled value-of-information test. The electron-only comparison measures whether classifier complexity by itself adds value. The electron-plus-t comparison measures whether independent nuclear-recoil information creates a nonlinear separation that the BDT can use more effectively than rectangular cuts. The exact-t calculation is deliberately an oracle test: it asks what would be possible if t were known perfectly. It does not assume that the EIC will reconstruct t with perfect acceptance or resolution.")
    paragraph(doc, "The central supported claim is correspondingly narrow. On our weighted parton-level signal and coherent-photon-proxy samples after the published Table I preselection, recoil-electron BDTs are practically equivalent to optimized rectangular cuts. When exact t is added to both methods, the BDT obtains a small but statistically robust advantage at 10 GeV. We do not claim that ML has already improved detector-level EIC sensitivity, that t is directly measured without uncertainty, that the public backgrounds have been fully reproduced, or that ML makes event generation faster.")

    heading(doc, "2. Physics Model and Classical Baseline", 1)
    heading(doc, "2.1. Coherent exclusive production process", 2)
    paragraph(doc, "We consider scalar and vector bosons coupled to electrons with benchmark coupling g_phi e = 10^-4 and an invisible branching fraction treated as unity. The signal process is e(k) + Au(P) -> e(k') + Au(P') + phi(p_phi), where the outgoing gold ion remains intact. The incoming electron energy is 18 GeV and the incoming gold energy is 100 GeV per nucleon. Gold is modeled with A = 197 and Z = 79, and a Helm elastic form factor describes the coherent nuclear response.")
    add_figure(doc, ROOT / "plots" / "exclusive_dark_boson_production.png", 1, "Coherent exclusive e(k) + Au(P) -> e(k') + Au(P') + phi(p_phi) production. The dark boson is emitted before panel (a) or after panel (b) virtual-photon exchange with intact Au. The nuclear vertex is Z F_A(t), with t = -(P' - P)^2, and the amplitude coherently sums both diagrams. The dashed line denotes a scalar or vector phi; the vector calculation uses the complete polarization sum.", width=5.45)
    paragraph(doc, "Figure 1 identifies the two leading-order electron-line emission diagrams. In panel (a), phi is radiated before the electron exchanges a virtual photon with the gold nucleus. In panel (b), phi is radiated after that exchange. Neither diagram is optional: the matrix element is their coherent sum, including interference. The double ion line emphasizes that the nucleus remains intact and couples through the elastic nuclear form factor.")
    paragraph(doc, "This mechanism must not be confused with exclusive vector-meson production. Exclusive vector-meson production normally refers to photon-hadron dynamics that create a hadronic vector meson, often described through gluon or Pomeron exchange. Here the scalar or vector dark boson is radiated directly from the electron line, while the nucleus supplies the coherent electromagnetic field. The word vector describes one possible spin assignment for phi; it does not turn the process into vector-meson production.")

    heading(doc, "2.2. Kinematic definitions and the meaning of t", 2)
    paragraph(doc, "All cuts and classifier inputs are evaluated in the collider laboratory frame. We choose the incoming gold beam to define +z and the incoming electron beam to travel in the opposite direction. Let q_e = k - k' be the four-momentum lost by the electron. The positive electron momentum-transfer variable is Q² = -q_e². The electron observables used by the reference selection are Q², pT,e, eta_e, and E_e.")
    paragraph(doc, "The recoil momentum-transfer variable is defined directly as t = -(P' - P)² > 0. This positive-t convention is convenient for spacelike exchange. Every equation, cut, and plot in this manuscript uses this same quantity, so no separate recoil-transfer symbol or opposite-sign convention is introduced.")
    paragraph(doc, "The variable t is already part of the physics calculation because it enters the coherent form factor and the three-body kinematics. What is new in the ML extension is not inserting t into the matrix element; it is giving the event-by-event value of t to both selectors as an additional discriminating input. In an ideal generated event, t is known exactly from P and P'. In a real experiment, it must be reconstructed from the outgoing ion or from a constrained event reconstruction and will have finite acceptance, efficiency, and resolution. Exclusivity makes t physically meaningful, but it does not guarantee that every outgoing ion is measurable.")

    heading(doc, "2.3. Event generation, weights, and cut efficiencies", 2)
    paragraph(doc, "Signal events are generated by direct numerical integration of the exact three-body phase space and the cited scalar and vector matrix elements [1,3,4]. The calculation includes the Helm form factor, the two electron-line emission amplitudes, their interference, and the full vector-polarization sum. Generated four-momenta are transformed to the collider laboratory frame before Q², pT,e, eta_e, E_e, t, and the electron-boson invariant mass are calculated.")
    paragraph(doc, "TFoam performs adaptive importance sampling, so generated events do not generally represent equal portions of the cross section. Each event carries a Monte Carlo weight. Cross sections and efficiencies are therefore calculated from sums of weights rather than from the fraction of stored entries that pass. For a selection C, the efficiency is epsilon_C = sum_i(w_i I_i(C))/sum_i(w_i), where I_i(C) equals one when event i passes and zero otherwise. Using unweighted event counts would bias the acceptance whenever event weights vary across phase space.")
    paragraph(doc, "The inclusive and Table I selected cross sections are obtained from the same physical matrix element and phase-space measure; no independent normalization is fitted to the selected curve. For the subsequent ML comparison, samples are generated conditional on the appropriate Table I preselection for each mass. The BDT and optimized rectangular selector therefore act only as second-stage refinements. They cannot recover signal events rejected by Table I, loosen the original preselection, or demonstrate faster event generation.")

    heading(doc, "2.4. Vector and scalar signal cross sections", 2)
    paragraph(doc, "Figure 2 presents our coherent production cross sections for invisible vector and scalar bosons over the mass range 0.01 to 10 GeV. We show the inclusive rates and the rates after the mass-dependent electron preselection described in Section 2.3, using the benchmark coupling g_phi e = 10^-4 and the beam configuration specified in Section 2.1.")
    add_figure(doc, REVISION_PLOTS / "vector_scalar_signal_cross_sections.png", 2, "Coherent vector and scalar signal cross sections as functions of boson mass for g_phi e = 10^-4 at 18 GeV x 100 GeV per nucleon. Green and purple show our inclusive and selected calculations, respectively. Solid lines denote vector bosons and dashed lines denote scalar bosons. The selected curves apply the mass-dependent electron preselection described in Section 2.3. Both axes are logarithmic.")
    paragraph(doc, "The inclusive cross sections decrease rapidly with increasing boson mass, and the vector rate exceeds the scalar rate throughout the scan. At m_phi = 1 GeV, our inclusive cross sections are 13.5868 pb for the vector and 5.96175 pb for the scalar. The logarithmic axes display the mass dependence across several orders of magnitude.")
    paragraph(doc, "The selected cross sections include the weighted efficiency of the electron preselection. At m_phi = 1 GeV, they are 0.157512 pb for the vector and 0.0372382 pb for the scalar, corresponding to efficiencies of approximately 1.16% and 0.625%, respectively. These selected signal samples provide the starting point for the subsequent comparison of BDTs and optimized rectangular cuts.")
    add_table(doc, ["Signal", "Inclusive [pb]", "Selected [pb]"], [
        ["Vector", "13.5868", "0.157512"],
        ["Scalar", "5.96175", "0.0372382"],
    ], [3120, 3120, 3120])
    add_caption(doc, "Table 1", "Inclusive and selected signal cross sections from our weighted numerical integration at m_phi = 1 GeV and g_phi e = 10^-4. The selected rates include the mass-dependent electron preselection.")

    heading(doc, "3. Background Model", 1)
    heading(doc, "3.1. Backgrounds treated in the reference analysis", 2)
    paragraph(doc, "The reference paper identifies coherent photon bremsstrahlung, e Au -> e Au gamma, as the dominant reducible background when the emitted photon is not detected. Its signal-like topology contains a recoil electron, an intact ion, and apparent missing energy. The paper assigns a miss probability of 10^-6 to central photons and assumes that photons outside |eta_gamma| = 3.5 are missed. These probabilities describe detector failure parametrically; they are not obtained from a full detector simulation.")
    paragraph(doc, "The reference also evaluates deep-inelastic scattering, eA -> eXj, at parton level with MadGraph5_aMC@NLO and assumes that a zero-degree calorimeter veto rejects 95% of the relevant nuclear-breakup background. Its Table I lists effective photon and DIS cross sections after each mass-dependent electron selection. An irreducible neutrino-pair background is reported to be negligible [1]. The reference does not pass its signal or backgrounds through Geant4.")

    heading(doc, "3.2. Coherent-photon proxy used in this study", 2)
    paragraph(doc, "No public background event files or generator cards were found in the current arXiv record. The older ancillary package contains coarse one-dimensional photon histograms but no event-level joint distribution from which the correlations among Q², pT,e, eta_e, E_e, and t can be recovered. Because multivariate classifiers learn precisely these joint correlations, the published effective cross sections alone are insufficient to train a physically validated background classifier.")
    paragraph(doc, "We therefore construct a reproducible coherent-photon proxy from the massless-vector limit of our coherent matrix element, using electromagnetic coupling e = 0.3028221209 and a photon-mass regulator of 10^-6 GeV. The regulator is far below every analyzed kinematic scale and only stabilizes the numerical representation of the massless limit. For each dark-boson mass hypothesis, the photon sample is generated with the same Table I electron preselection used for the corresponding signal comparison.")
    paragraph(doc, "Each photon event is weighted by the Monte Carlo integration weight and the reference miss prescription. Central photons receive a miss factor of 10^-6, while photons with |eta_gamma| > 3.5 receive unit miss probability. The resulting shape is used to compare algorithms, while the absolute photon rate entering the significance proxy is normalized to the corresponding effective Table I value. This procedure prevents a raw generator-rate mismatch from trivially determining the comparison, but it does not repair an incorrect multivariate shape.")
    paragraph(doc, "As an external diagnostic rather than a tuning target, we compare the inclusive photon proxy with the older public one-dimensional histograms in Figure 3. Because the public files provide only the displayed bins and no event count or uncertainty, each curve is normalized independently within the published range. The resulting range-conditioned total-variation distances are 0.715 for log10(pT,e) and 0.409 for eta_e. In addition, 37.3% of our weighted pT,e distribution and 50.3% of our weighted eta_e distribution lie outside the public plotting ranges. These facts prevent the comparison from being interpreted as full-distribution agreement.")
    add_figure(doc, REVISION_PLOTS / "photon_shape_overlay.png", 3, "Inclusive photon-shape diagnostic using the binning in the older public ancillary distributions. The left and right panels show log10(pT,e/GeV) and eta_e. Black denotes the public coherent-photon marginal and blue denotes our massless-vector photon proxy. Each in-range curve is normalized independently, so the quoted total-variation distances of 0.715 and 0.409 are range-conditioned. The annotations also report that 37.3% and 50.3% of our weighted proxy lie outside the respective public ranges. No public uncertainties or event-level correlations are available. This diagnostic is not used to reweight the proxy or train either selector.", width=6.45)
    paragraph(doc, "Figure 3 is therefore a transparent failure of marginal shape reproduction, not a successful validation plot. We do not reweight to the coarse histograms because separate one-dimensional distributions cannot determine the correlated rare tail surviving Table I, and artificial marginal agreement could create unjustified joint correlations. The photon shape is the dominant unresolved model limitation and directly conditions every ML-versus-cuts statement in this paper.")

    heading(doc, "3.3. DIS treatment and interpretation boundary", 2)
    paragraph(doc, "A separate rate-level cross-check was completed with MadGraph5_aMC@NLO 3.7.0 for leading-order ep -> ej using CT10nlo, scaled by A = 197, 5% veto survival, and the stated 10^-6 jet-miss factor. Across the 11 Table I mass-dependent selections, our effective DIS estimates differ from the published values by 4.30% to 6.23%. The analogous coherent-photon proxy rates are larger than the published photon estimates by factors of 3.64 to 4.11. Thus the DIS normalization is reproduced at the stated rate level, while the photon-rate and photon-shape discrepancies remain unresolved.")
    paragraph(doc, "No event-level DIS sample is available in this project. We retain the effective DIS rate reported in Table I but set its second-stage selection efficiency to unity for both methods. In other words, neither the BDT nor the optimized rectangular selector is credited with rejecting any additional DIS events. This makes each quoted improvement relative to the unchanged Table I selection smaller than it would be in a photon-only calculation.")
    paragraph(doc, "Leaving DIS unrejected does not guarantee that the BDT-versus-cuts ordering is conservative in the real experiment. The true BDT and cut efficiencies for DIS could differ, and that unknown difference could increase, reduce, or reverse the small advantage reported here. Consequently, our numerical results are conditional algorithm comparisons within a stated photon proxy and an unrejected DIS component. They are not validated predictions of total EIC background rejection or absolute coupling reach.")

    heading(doc, "4. Fair Machine-Learning Test", 1)
    heading(doc, "4.1. Two nested information sets", 2)
    paragraph(doc, "The primary, electron-only input set contains log10(Q²), pT,e, eta_e, and E_e. These quantities correspond to measurements that can in principle be reconstructed from the outgoing electron. In the present study they are evaluated exactly from generator-level four-momenta: no tracking efficiency, energy scale, angular resolution, bremsstrahlung recovery, acceptance loss, or variable-to-variable detector correlation is applied. At fixed beam energy, pT,e and eta_e already determine most of the outgoing-electron four-momentum, so E_e and Q² are substantially redundant. This makes the electron-only test intentionally demanding for ML: the model receives several representations of roughly two independent electron degrees of freedom.")
    paragraph(doc, "The extended information set contains the same four electron quantities plus exact generator-level t. This is a nested comparison; no electron variable is removed or changed. Crucially, both the BDT and the rectangular selector receive t in the blue comparison. The blue result is therefore not the gain from giving a privileged variable only to ML. It is the residual gain from using a nonlinear BDT boundary after both algorithms receive the additional recoil information.")
    paragraph(doc, "The exact value of t is labeled an oracle input throughout the paper. It represents the best-case information available if the outgoing ion momentum were known perfectly. It is not called a detector observable, and the result is not interpreted as a measured EIC reach. The purpose of the oracle is to determine whether a realistic reconstructed-t study is scientifically motivated before investing in a detector model.")

    heading(doc, "4.2. Identical data roles and prevention of test leakage", 2)
    paragraph(doc, "For every signal type, boson mass, and input set, the BDT and rectangular selector use the same weighted signal and photon events. Independently generated samples have distinct roles. The training sample is used to fit the BDT and construct rectangular candidates. The validation sample selects the BDT profile, BDT score threshold, rectangular grid, and final box boundaries. These choices are frozen before the held-out test sample is evaluated.")
    paragraph(doc, "The held-out test events are used once for the reported comparison. The two methods are evaluated on the same test events, so their difference can be estimated with a paired bootstrap instead of treating the two results as statistically independent. Source files, executables, configurations, selected models, cut boxes, generated data, and final results are recorded in checksum manifests. This design limits accidental leakage, post-test retuning, and comparisons between mismatched generator versions.")

    heading(doc, "4.3. BDT and rectangular-selector definitions", 2)
    paragraph(doc, "The ML method is a boosted decision tree implemented with TMVA. At each of the 11 Table I masses for both vector and scalar signals, four prespecified model profiles are considered: 400-tree models with maximum depths two, three, and four, and an 800-tree depth-three model. Training fits the trees, and validation selects among these fixed profiles and chooses the score threshold that maximizes the common analysis objective.")
    paragraph(doc, "The classical comparator is not merely the unchanged Table I selection. It is a weighted multistart rectangular selector using exactly the same input variables as the BDT. Candidate lower and upper bounds are generated from weighted-threshold grids containing 40, 80, and 160 quantiles, with multiple starting points to reduce dependence on a single greedy solution. Training constructs candidate boxes and validation selects the final grid and bounds. The method is a strong approximate rectangular baseline, although it is not claimed to be the mathematically global optimum over every possible box.")
    paragraph(doc, "Using the same inputs is more important than giving the two algorithms the same formal complexity. A BDT with hundreds of trees is inherently more flexible than a rectangular box. The fair question is whether that extra flexibility improves the final metric when samples, weights, information, preselection, and data roles are controlled. Calling the methods 'same complexity' would be incorrect; we instead call them a same-information comparison.")

    heading(doc, "4.4. Optimization target, uncertainty, and decision rules", 2)
    paragraph(doc, "Both methods act only as second-stage refinements after the published Table I preselection. They maximize the same approximate background-dominated significance ratio")
    equation(doc, "R_Z = ε_S √[(σ_γ + σ_DIS)/(ε_γ σ_γ + σ_DIS)].")
    paragraph(doc, "Here epsilon_S and epsilon_gamma are the weighted signal and photon efficiencies of the second-stage selector, while sigma_gamma and sigma_DIS are the effective Table I background cross sections. The DIS term is left unrejected. R_Z = 1 means that the second-stage selector provides no improvement relative to stopping after Table I; values above one indicate improvement within this proxy. The direct algorithm comparison is the ratio R_Z(BDT)/R_Z(rectangular).")
    equation(doc, "A_g = 1 - √[R_Z(rectangular)/R_Z(BDT)].")
    paragraph(doc, "Because the signal cross section scales as the square of the electron-boson coupling in this model, A_g is the relative reduction in the coupling threshold implied by the ratio of the two approximate significance metrics. Positive values favor the BDT and negative values favor rectangular cuts. This quantity is an algorithmic coupling-threshold proxy, not an absolute exclusion or discovery reach.")
    paragraph(doc, "A +1% coupling advantage was fixed before the final robustness evaluation as the minimum practically useful effect. This is a project decision threshold, not a universal EIC or community standard. Through the equation above, a 1% coupling advantage corresponds to approximately a 2.03% increase in R_Z over the rectangular selector. The threshold prevents a statistically resolvable but scientifically negligible sub-percent fluctuation from being described as a useful ML improvement.")
    paragraph(doc, "The full electron-only mass scan contains 22 primary comparisons: 11 masses for each of two signal spins. Paired event bootstraps provide pointwise intervals, and a Bonferroni lower bound controls the family of primary tests. A separate robustness study repeats the complete event generation, training, validation, model selection, cut optimization, and sealed-test evaluation ten times at four prespecified electron-only benchmarks: vector masses 1 and 10 GeV and scalar masses 1 and 6.31 GeV. It contains 40 complete replicas, 2,000 paired event-bootstrap samples per replica, and 20,000 hierarchical bootstrap draws per benchmark.")
    paragraph(doc, "The exact-t robustness study repeats the same protocol at the two positive high-mass benchmarks: vector 10 GeV and scalar 10 GeV. It contains ten complete replicas per signal, 2,000 paired event-bootstrap samples per replica, and 20,000 hierarchical bootstrap draws per signal. Every replica regenerates independent training, validation, and test events, refits all four BDT profiles, and reoptimizes all three rectangular grid settings. The BDT and cuts receive identical Q², pT,e, eta_e, E_e, and exact-t inputs. All method and result checksum manifests pass verification.")

    heading(doc, "5. Results", 1)
    heading(doc, "5.1. Weighted kinematic structure after Table I", 2)
    paragraph(doc, "Before comparing classifiers, Figure 4 shows the weighted joint structure that supplies the additional information. The horizontal coordinate is the recoil-electron transverse momentum, and the vertical coordinate is log10(t/GeV²). Rows correspond to 1 and 10 GeV; columns show the coherent-photon proxy, vector signal, and scalar signal. Each panel is normalized to its own total weighted class probability, so color shows shape rather than the very different absolute rates. The photon weights include the reference miss prescription used in the classifier study.")
    add_figure(doc, REVISION_PLOTS / "t_pt_density.png", 4, "Weighted post-Table-I distributions in log10(t/GeV²) and pT,e. Rows correspond to m_phi = 1 and 10 GeV; columns show the coherent-photon proxy, vector signal, and scalar signal. Each panel is normalized to its total weighted class probability, so color represents shape rather than absolute cross section. Signal events use their integration weights, while photon events additionally include the stated photon-miss probability. Identical axes, bins, and color scales are used within each mass row. The value of t is exact generator truth and is shown as an oracle input, not as a detector-level reconstructed observable.", width=6.65)
    paragraph(doc, "At 1 GeV, all three classes populate broadly similar wedge-shaped regions in the t-pT,e plane. At 10 GeV, the photon proxy retains substantially more weight at smaller t, while both signal hypotheses shift toward larger t and exhibit a different correlation with pT,e. A single independent interval in each variable cannot follow the complete sloped boundary. This visual pattern is consistent with, but does not by itself prove, the later finding that exact t is useful to the nonlinear BDT mainly at high mass.")

    heading(doc, "5.2. Impact of exact t across boson mass", 2)
    paragraph(doc, "Figures 5 and 6 provide the direct answer to the main algorithmic question for vector and scalar signals. The horizontal axis is the assumed boson mass on a logarithmic scale. The vertical axis is A_g, the BDT coupling-threshold advantage over the same-input optimized rectangular selector, expressed as a percentage. A value of zero means equal performance. Positive values favor the BDT; negative values favor rectangular cuts. The horizontal dotted line at +1% is the predefined practical-usefulness threshold. Error bars are paired 95% event-bootstrap intervals for the fixed trained model and cut box in the sealed mass scan.")
    add_figure(doc, ROOT / "plots" / "ml_fair_vector_t_comparison.png", 5, "Vector mass scan. The logarithmic horizontal axis is the vector-boson mass. The vertical axis is the BDT coupling-threshold advantage over same-input optimized rectangular cuts; zero denotes equal performance and positive values favor the BDT. Solid black circles use Q², pT,e, eta_e, and E_e. Solid blue circles give both algorithms the same four variables plus exact generator-level t. Error bars are paired 95% event-bootstrap intervals. The dotted horizontal line marks the predefined +1% practical threshold.")
    paragraph(doc, "For the vector signal, the black electron-only points remain close to zero at every mass. The BDT therefore finds no useful nonlinear separation beyond the optimized box when it is restricted to the recoil-electron information. The blue exact-t points also remain near zero at low mass, but they rise above zero as the mass increases and reach approximately +1.14% at 10 GeV in the single sealed scan. This pattern suggests that the relevant nonlinear signal-background correlation is associated with the nuclear recoil and becomes most useful near the upper end of the studied vector-mass range.")
    add_figure(doc, ROOT / "plots" / "ml_fair_scalar_t_comparison.png", 6, "Scalar mass scan. Axes and statistical conventions match Figure 5. Solid black circles use the four recoil-electron inputs. Solid blue circles add exact generator-level t to both the BDT and optimized rectangular selector. The comparison therefore measures the BDT's nonlinear advantage after both methods receive t; it does not compare an ML method with t against cuts without t. The blue curve is an oracle result, not a detector-level projection.")
    paragraph(doc, "The scalar scan leads to the same qualitative conclusion. The electron-only black points fluctuate close to zero, including small positive and negative differences. Adding exact t produces a clearer positive trend above approximately 2 GeV and reaches +1.37% at 10 GeV in the sealed mass scan. The vector and scalar curves therefore support a shared interpretation: exact nuclear-recoil information is most valuable to the BDT at high mass, while electron-only classifier complexity is not sufficient.")
    paragraph(doc, "Across all 22 electron-only scalar/vector points, the nominal coupling-threshold difference ranges from -0.228% to +0.161%. No point has a positive familywise-controlled lower bound, and no point reaches the +1% practical threshold. The largest nominal positive electron-only result occurs for the 1.585 GeV scalar: +0.161%, with a pointwise 95% interval from +0.033% to +0.289%, but its familywise lower bound is -0.026%. Several points slightly favor rectangular cuts. The complete scan therefore provides no multiplicity-controlled evidence that an electron-only BDT is better.")
    paragraph(doc, "This conclusion does not mean that the Table I selection is already optimal. Both second-stage methods can improve the scalar proxy relative to leaving Table I unchanged. Around 1--2.5 GeV, their R_Z values reach approximately 1.05--1.11. The important result is that the optimized rectangular selector captures essentially all of this electron-only improvement. The gain is due to reoptimizing the selection on our samples, not specifically due to machine learning.")
    paragraph(doc, "The exact-t curves answer a different question. In the single sealed mass scan at 10 GeV, the BDT advantage is +1.14% with a pointwise 95% interval of [+0.96%, +1.32%] for the vector and +1.37% with an interval of [+1.15%, +1.59%] for the scalar. These single-sample intervals alone would not demonstrate that the effect survives new event generation and retraining. Section 5.4 therefore repeats the entire analysis across independent seeds.")

    heading(doc, "5.3. Operating characteristics and variable attribution at 10 GeV", 2)
    paragraph(doc, "Figure 7 separates the two ingredients of the final comparison. A continuous BDT curve is obtained by scanning the frozen score threshold on the sealed post-Table-I test events. Each optimized rectangular selector supplies one validation-selected operating point rather than a continuous ROC curve. The horizontal axis is the weighted survival efficiency of the coherent-photon proxy, epsilon_gamma, and the vertical axis is weighted signal efficiency, epsilon_S. Curves closer to the upper-left corner rank signal above photon more effectively. The plot contains no event-level DIS because no such sample is classified in this study.")
    add_figure(doc, REVISION_PLOTS / "roc_t_comparison.png", 7, "Weighted post-Table-I signal efficiency epsilon_S versus coherent-photon-proxy efficiency epsilon_gamma on the sealed 10 GeV test samples. The left and right panels show vector and scalar signals. Solid curves scan the frozen BDT score threshold; filled circles are the validation-selected optimized rectangular operating points. Black denotes Q², pT,e, eta_e, and E_e, while blue adds exact generator-level t to both methods. The weighted sealed AUC values are 0.674 and 0.822 for the vector and 0.721 and 0.839 for the scalar without and with t, respectively. DIS is absent from this diagnostic because no event-level DIS sample is classified.", width=6.55)
    paragraph(doc, "For the vector signal, exact t raises the sealed weighted AUC from 0.674 to 0.822; for the scalar it raises the AUC from 0.721 to 0.839. The blue curves therefore show a substantial improvement in signal-photon ranking across thresholds, even though the final coupling-threshold advantage over a same-input optimized box is only 1.1% to 1.4%. AUC and coupling advantage answer different questions: AUC integrates ranking performance over every threshold, whereas A_g compares two validation-selected working points after the unrejected DIS rate is included in R_Z.")
    paragraph(doc, "Figure 8 provides a model-level check of which inputs the exact-t BDTs actually use. For each of the ten independent 10 GeV generation-and-training replicas, the validation-selected TMVA model reports a normalized method-specific variable importance. The plot shows every replica together with the mean and sample standard deviation for each feature.")
    add_figure(doc, REVISION_PLOTS / "tmva_t_importance.png", 8, "TMVA BDT variable importance for the validation-selected exact-t models across ten independent 10 GeV generation-and-training replicas. Green and purple denote vector and scalar signals. Small points show individual replicas; large circles and error bars show the mean and sample standard deviation. Importance is normalized within each fitted model. Exact t is the highest-ranked input for both signals, with mean importances 0.2786 +/- 0.0143 for the vector and 0.2920 +/- 0.0148 for the scalar. These training-derived, method-specific rankings are affected by correlations among inputs and are not causal or detector-level feature attributions.", width=6.15)
    paragraph(doc, "Exact t is ranked highest in both ensembles. Its mean importance is 0.2786 +/- 0.0143 for the vector and 0.2920 +/- 0.0148 for the scalar. The remaining mean importances span 0.164 to 0.212. This ranking is consistent with the density and ROC plots: the positive high-mass result is associated with additional recoil information rather than with a different electron-only model choice. However, TMVA importance is training-derived and correlation-dependent. It neither proves causation nor shows that a realistically reconstructed t would retain the same value.")

    heading(doc, "5.4. Independent-seed robustness", 2)
    paragraph(doc, "A single train/validation/test split can give a stable-looking answer that depends on one adaptive Monte Carlo sample or one training seed. The robustness studies therefore repeat the complete workflow, not merely the final event bootstrap. Each replica regenerates signal and photon samples, retrains the candidate BDTs, reoptimizes rectangular boxes, freezes both methods on validation data, and evaluates the pair on new held-out test events.")
    add_figure(doc, ROOT / "plots" / "ml_seed_robustness.png", 9, "Ten-replica electron-only robustness comparison. Each row is one prespecified signal benchmark. The horizontal axis is the BDT coupling-threshold advantage over same-input optimized rectangular cuts; zero means equal performance. Each point is the hierarchical median over ten complete generation-and-training replicas, and each horizontal bar is the corresponding 95% interval. All four intervals lie far inside the predefined -1% to +1% practical-equivalence region.")
    add_table(doc, ["Signal", "m_phi [GeV]", "BDT coupling advantage [95%]", "BDT wins", "Decision"], [
        ["Vector", "1.000", "-0.009% [-0.029%, +0.004%]", "4/10", "Practical equivalence"],
        ["Vector", "10.000", "-0.057% [-0.142%, +0.021%]", "2/10", "Practical equivalence"],
        ["Scalar", "1.000", "+0.031% [-0.083%, +0.132%]", "8/10", "Practical equivalence"],
        ["Scalar", "6.310", "-0.088% [-0.182%, -0.005%]", "0/10", "Cuts slightly favored"],
    ], [1500, 1300, 3000, 1000, 2560])
    add_caption(doc, "Table 2", "Independent-seed results. All four ordinary 95% intervals lie inside the preregistered -1% to +1% practical-equivalence range; none passes the familywise superiority test.")
    paragraph(doc, "Figure 9 and Table 2 show that the electron-only conclusion is not caused by one unlucky model or sample. Vector 1 GeV, vector 10 GeV, scalar 1 GeV, and scalar 6.31 GeV all fall within the preregistered practical-equivalence margin. No benchmark has a familywise lower bound above zero, and none approaches a +1% BDT advantage. Scalar 6.31 GeV has a 95% interval slightly below zero, so optimized cuts are favored statistically in that isolated comparison, but the magnitude is only about one tenth of one percent and remains scientifically negligible under the predefined margin.")
    paragraph(doc, "Selected BDT profiles and rectangular grid sizes vary among replicas because their validation scores are nearly degenerate, yet the final performance ratio remains near one. This is useful evidence: the conclusion is stable even though the nominally selected implementation changes. The electron-only result is therefore practical equivalence, not a hidden positive ML result obscured by one architecture choice.")
    add_figure(doc, ROOT / "plots" / "ml_t_seed_robustness.png", 10, "Ten-replica exact-t robustness comparison at 10 GeV. The horizontal axis is the BDT coupling-threshold advantage over same-input optimized cuts. Blue circles are hierarchical medians and horizontal bars are 95% intervals over complete independent generation-and-training replicas. The dotted vertical line marks the predefined +1% practical threshold. Both methods receive Q², pT,e, eta_e, E_e, and exact generator-level t.")
    add_table(doc, ["Signal", "m_phi [GeV]", "BDT advantage [95%]", "BDT wins", "Familywise lower", "Decision"], [
        ["Vector", "10.000", "+1.208% [+1.081%, +1.326%]", "10/10", "+1.081%", "Practical BDT superiority"],
        ["Scalar", "10.000", "+1.361% [+1.206%, +1.529%]", "10/10", "+1.206%", "Practical BDT superiority"],
    ], [1150, 1050, 2800, 900, 1500, 1960])
    add_caption(doc, "Table 3", "Independent-seed exact-t results. Hierarchical paired intervals combine complete generator-and-training replicas with paired event resampling. Both familywise lower bounds exceed the predefined +1% practical threshold.")
    paragraph(doc, "Figure 10 and Table 3 show that the exact-t result survives every independent replica. The BDT wins 10 of 10 vector comparisons and 10 of 10 scalar comparisons. The hierarchical median coupling-threshold advantages are +1.208% [1.081%, 1.326%] for the vector and +1.361% [1.206%, 1.529%] for the scalar. Their two-test familywise lower bounds are +1.081% and +1.206%, respectively. Both lower bounds remain above the predefined +1% threshold. Within the stated generator-level model, this passes the statistical-superiority and practical-usefulness decision rules.")
    paragraph(doc, "The difference between Figures 9 and 10 is the key robustness result. Figure 9 shows that additional model complexity does not help when both methods receive only redundant electron kinematics. Figure 10 shows that the BDT becomes modestly but consistently better when both methods receive exact t. The effect is therefore attributable to how the BDT combines new recoil information with the electron variables, not to an unfair input advantage or a comparison against the unreoptimized Table I box.")

    heading(doc, "5.5. Size and scientific meaning of the high-mass effect", 2)
    paragraph(doc, "The observed 1.2--1.4% coupling-threshold advantages are numerically small. They should not be described as a dramatic increase in sensitivity. However, they are not the same as a 1.2--1.4% increase in the underlying significance proxy. Inverting the coupling transformation, the vector median corresponds to R_Z(BDT)/R_Z(rectangular) of approximately 1.0246, while the scalar median corresponds to approximately 1.0278. Thus the BDT improves the stated significance proxy by roughly 2.5% and 2.8% relative to optimized cuts in the exact-t 10 GeV study.")
    paragraph(doc, "The importance of the result is not its size alone. The effect was tested after both methods received identical inputs, after a strong rectangular reoptimization, on held-out weighted samples, and across ten independently regenerated and retrained replicas. The confidence intervals exclude zero and their familywise lower bounds exceed the threshold selected before the robustness result was opened. The result therefore establishes a reproducible value-of-information effect inside the stated model.")
    paragraph(doc, "At the same time, statistical robustness is not equivalent to experimental validity. The intervals quantify finite simulated-event variation and observed generation/training-seed variation. They do not contain uncertainty from detector acceptance, reconstructed-t resolution, photon-shape modeling, DIS selection, luminosity, nuclear modeling, or the scalar matrix element. A 1.2--1.4% oracle advantage could shrink, disappear, or change after those effects are introduced. The correct publication claim is therefore that exact nuclear-recoil information creates a small and robust generator-level ML advantage at 10 GeV, motivating a reconstructed-t study; it is not yet a claim of improved experimental reach.")

    heading(doc, "6. Discussion and Next Stage", 1)
    heading(doc, "6.1. Why t changes the machine-learning comparison", 2)
    paragraph(doc, "The two input sets give an information-based explanation for the otherwise different conclusions. The electron-only variables all describe one recoil electron. At fixed incoming energy, pT,e and eta_e determine the electron direction and most of its energy-momentum, while E_e and Q² provide correlated representations of the same object. Once the rectangular selector is reoptimized on the same weighted samples, it captures essentially all available separation. A BDT can represent more complicated boundaries, but it cannot create information that is absent from its inputs.")
    paragraph(doc, "Exact t adds a variable from the intact-ion side of the exclusive event. Because phi is invisible, the electron alone does not completely determine the recoil configuration. The relationship among t and the electron quantities changes across signal and coherent-photon phase space, especially toward high boson mass. A BDT can partition this multidimensional space into several correlated regions, whereas a rectangular selector can only impose independent upper and lower bounds on each variable. The 10 GeV results indicate that this nonlinear combination leaves a small residual advantage after the box has been strongly optimized.")
    paragraph(doc, "This interpretation does not imply that t is always more useful at higher mass under every detector or background model. It describes the trend in the current generator-level samples. Detector smearing may erase the relevant correlation, while reconstructed recoil, veto response, shower shapes, track quality, missing momentum, and forward-detector information may introduce additional correlations. The exact-t result supplies a concrete hypothesis: ML becomes useful when reconstruction retains independent recoil information with enough precision that nonlinear signal-background structure survives.")

    heading(doc, "6.2. Physics limitations", 2)
    paragraph(doc, "The signal rates presented in Section 2.4 are obtained from our scalar and vector implementations. The following limitations concern the background model, detector response, and statistical interpretation of the signal-selection study.")
    paragraph(doc, "The dominant limitation is the background model. Within the displayed public ranges, the coherent-photon proxy does not reproduce the available one-dimensional marginals: the range-conditioned total-variation distances are 0.715 in log10(pT,e) and 0.409 in eta_e, while substantial proxy weight also lies outside those ranges. At the rate level, our photon estimate exceeds the published Table I values by factors of 3.64 to 4.11. More importantly, the joint selected-tail distribution learned by the BDT is not publicly available for validation. Renormalizing the proxy to a Table I cross section controls only its total rate after preselection; it cannot validate the multivariate shape that determines BDT-versus-cuts performance.")
    paragraph(doc, "The independent DIS rate cross-check agrees with the published effective values within 4.30% to 6.23%, but no event-level DIS sample is classified. Leaving the published DIS rate unrejected is conservative relative to the unrefined Table I result for either selector separately, but it does not prove that the real BDT-versus-cuts ordering is conservative because the two unknown DIS efficiencies may differ. Rare photon-miss probabilities near 10^-6 also cannot be validated by ordinary samples containing only thousands or millions of events without a detector-tail model or control-region strategy.")
    paragraph(doc, "No detector response is applied. The electron variables are exact generator-level quantities, and t is calculated from the exact outgoing-ion four-momentum. The analysis contains no far-forward ion acceptance, beam optics, reconstruction efficiency, t smearing, photon-veto response, zero-degree-calorimeter resolution, pileup, beam background, or detector-induced correlation. The exact-t result should therefore be read as an upper-bound value-of-information experiment.")
    paragraph(doc, "Finally, R_Z is a background-dominated S/sqrt(B)-type proxy rather than a nuisance-aware profile likelihood. The paired and hierarchical intervals quantify finite simulated-event and observed generation/training-seed variation. They do not include detector, generator, luminosity, nuclear-model, or background-normalization systematics. These limitations prevent an absolute coupling-reach or detector-level EIC sensitivity claim.")

    heading(doc, "6.3. From exact t to future EIC expectations", 2)
    paragraph(doc, "The next stage should convert exact t into a reconstructed quantity through a sequence of controlled tests. First, the scalar matrix element should receive an independent analytic or separate-generator validation. The coherent-photon calculation should be compared with an independent implementation, and event-level DIS samples should be generated. Where exact agreement is impossible, physically motivated generator and rate variations should bracket the uncertainty instead of being hidden by a single normalization.")
    paragraph(doc, "Second, a documented fast detector model should be introduced. It should smear and apply efficiencies to the recoil electron, model whether the intact outgoing ion remains within the far-forward acceptance, and reconstruct t from the measured ion momentum when possible. The analysis should report t acceptance, efficiency, bias, and resolution as functions of mass and true t. Photon-veto and zero-degree-calorimeter responses should be included because the result depends on missed photons and rejected nuclear breakup.")
    paragraph(doc, "Third, both methods must receive the same reconstructed inputs. The correct detector comparison is reconstructed-electron plus reconstructed-t BDT versus reconstructed-electron plus reconstructed-t optimized cuts. Giving reconstructed t only to ML would confound new information with algorithm choice. Exact t should remain in plots as an oracle upper bound, while an electron-only detector curve provides the lower-information reference.")
    paragraph(doc, "Fourth, the complete training, validation, sealed-test, and independent-seed protocol should be repeated after detector effects are frozen. A nuisance-aware profile likelihood should replace R_Z for sensitivity statements and should include photon, DIS, detector, luminosity, and generator nuisance parameters. Only then should the work quote absolute expected coupling reach or claim an experimentally meaningful EIC improvement.")
    paragraph(doc, "A full Geant4 or ePIC simulation is not required for the first detector study and was not used by the reference analysis. A transparent fast-smearing and efficiency model is the appropriate first falsification test. If the t-driven advantage disappears immediately, the oracle result has still identified why. If it survives reasonable detector and background variations, selected benchmark points can then be validated with a fuller simulation.")

    heading(doc, "6.4. What the present paper supports", 2)
    paragraph(doc, "Supported statement. On the specified weighted parton-level signal and coherent-photon-proxy samples, conditional on the mass-dependent Table I preselection and published effective background normalizations, recoil-electron BDTs are practically equivalent to strong optimized rectangular cuts. When exact generator-level t is supplied to both methods, the BDT shows a small, repeatable 10 GeV advantage for both signal spins, with ten of ten wins and familywise lower bounds above the predefined +1% coupling threshold.", bold_lead="Supported statement. ")
    paragraph(doc, "Physical interpretation. The comparison demonstrates that the useful ingredient is additional independent recoil information, not ML complexity alone. The exact-t result is evidence that a reconstructed nuclear-recoil study is worth performing and provides a quantitative upper-bound benchmark against which detector degradation can be measured.", bold_lead="Physical interpretation. ")
    paragraph(doc, "Unsupported statement. The current work does not establish that ML improves real EIC sensitivity, does not show that t is reconstructed perfectly or even accepted for every exclusive event, does not validate the full photon and DIS backgrounds, does not replace the Table I preselection, does not demonstrate faster event production, and does not provide a detector-level coupling reach. Those statements require the next-stage validation described above.", bold_lead="Unsupported statement. ")

    heading(doc, "7. Conclusions", 1)
    paragraph(doc, "We calculated coherent exclusive signal production e Au -> e Au phi for scalar and vector dark bosons and used these signal samples to perform a controlled comparison between boosted decision trees and optimized rectangular cuts. The production mechanism is electron-line dark-boson radiation in the coherent electromagnetic field of an intact gold nucleus. We presented the inclusive and selected cross sections across the boson-mass scan, with the selected samples providing the baseline for both algorithms.")
    paragraph(doc, "The electron-only test gives a clear negative result. When the BDT and the rectangular selector receive the same Q², pT,e, eta_e, and E_e information, the complete 22-point mass scan finds no familywise-controlled evidence of BDT superiority. Ten independent replicas at four benchmarks place all ordinary 95% intervals well inside the predefined -1% to +1% practical-equivalence range. Additional classifier complexity does not produce a useful gain from these redundant generator-level electron variables.")
    paragraph(doc, "The exact-t test changes the high-mass comparison. When both methods receive the same electron quantities plus exact generator-level t, the BDT obtains median coupling-threshold advantages of +1.208% for the 10 GeV vector and +1.361% for the 10 GeV scalar. It wins all ten independently generated and retrained replicas for both signals. The hierarchical 95% intervals are [1.081%, 1.326%] and [1.206%, 1.529%], and the familywise lower bounds remain above the predefined +1% practical threshold. These coupling differences correspond to approximately 2.5% and 2.8% improvements in the stated R_Z proxy relative to same-input optimized cuts.")
    paragraph(doc, "The correct takeaway is not that ML is universally superior or that a detector-level EIC reach has already improved. The result shows that additional independent physical information can matter more than classifier complexity. Exact t exposes a modest but robust nonlinear recoil correlation at high mass that is not captured fully by a rectangular box. This makes reconstructed nuclear recoil a motivated next input to test, while the exact-t curves provide an oracle upper-bound benchmark.")
    paragraph(doc, "The conclusion remains conditional on the present physics model. The coherent-photon joint shape is not independently validated, no event-level DIS sample is classified, the scalar amplitude lacks an independent implementation, and no detector response or profile likelihood is included. The immediate next step is therefore to compare detector-level BDTs and reoptimized cuts using identical reconstructed electron and reconstructed-t information. The 10 GeV advantage becomes an experimental result only if it survives realistic ion acceptance, t resolution, photon and DIS variations, systematic uncertainties, independent seeds, and likelihood-based inference.")

    heading(doc, "Author Contributions", 1)
    paragraph(doc, "Conceptualization, methodology, software, validation, formal analysis, visualization, data curation, and writing - original draft: Rojae Mighty.")
    heading(doc, "Acknowledgments", 1)
    paragraph(doc, "The author thanks Hongkai Liu for helpful discussions of coherent dark-boson production.")
    heading(doc, "Data Availability Statement", 1)
    paragraph(doc, "The accompanying project package contains source code, configuration tables, generated ROOT files, weighted held-out scores, trained TMVA XML models, full-scan, electron-only ten-seed, and exact-t ten-seed numerical results, bootstrap distributions, checksum manifests, background rate comparisons, figure-level CSV data, and plotting scripts. The published authors' background event files and generator cards were not available and are therefore not redistributed.")
    heading(doc, "Conflicts of Interest", 1)
    paragraph(doc, "The author declares no conflict of interest.")
    heading(doc, "References", 1)
    refs = [
        "1. Davoudiasl, H.; Liu, H. Electron-Ion Collider as a Discovery Tool for Invisible Dark Bosons. arXiv 2025, arXiv:2505.08871; revised 2026.",
        "2. Davoudiasl, H.; Liu, H. Erratum: Electron-Ion Collider as a Discovery Tool for Invisible Dark Bosons. Phys. Rev. D 2026, 114, 019902(E). https://doi.org/10.1103/nw5p-k61z.",
        "3. Davoudiasl, H.; Marcarelli, R.; Neil, E.T. Displaced Signals of Hidden Vectors at the Electron-Ion Collider. Phys. Rev. D 2023, 108, 075017. https://doi.org/10.1103/PhysRevD.108.075017.",
        "4. Balkin, R.; Hen, O.; Li, W.; Liu, H.; Ma, T.; Soreq, Y.; Williams, M. Probing Axion-Like Particles at the Electron-Ion Collider. JHEP 2024, 02, 123. https://doi.org/10.1007/JHEP02(2024)123.",
        "5. Hoecker, A.; Speckmayer, P.; Stelzer, J.; Therhaag, J.; von Toerne, E.; Voss, H. TMVA: Toolkit for Multivariate Data Analysis. PoS ACAT 2007, 040. arXiv:physics/0703039.",
        "6. Alwall, J.; Herquet, M.; Maltoni, F.; Mattelaer, O.; Stelzer, T. MadGraph 5: Going Beyond. JHEP 2011, 06, 128. https://doi.org/10.1007/JHEP06(2011)128.",
        "7. Abdul Khalek, R. et al. Science Requirements and Detector Concepts for the Electron-Ion Collider. Nucl. Phys. A 2022, 1026, 122447. https://doi.org/10.1016/j.nuclphysa.2022.122447.",
        "8. Helm, R.H. Inelastic and Elastic Scattering of 187-MeV Electrons from Selected Even-Even Nuclei. Phys. Rev. 1956, 104, 1466-1475. https://doi.org/10.1103/PhysRev.104.1466.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.02
        set_font(p.add_run(ref), size=8.7)

    core = doc.core_properties
    core.title = "When Does Nuclear-Recoil Information Make Machine Learning Useful for Invisible Dark-Boson Selection at the Electron-Ion Collider?"
    core.author = "Rojae Mighty"
    core.subject = "A parton-level value-of-information study"
    core.keywords = "EIC, dark bosons, machine learning, nuclear recoil, momentum transfer, optimized cuts, TMVA"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
